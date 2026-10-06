#!/usr/bin/env python3
"""Build the PPADEM Word reference document, template and example.

Starts from Pandoc's default reference.docx and restyles it with the brand
palette from _brand/ppadem-brand.scss. Writes:

    _extensions/ppadem-word/ppadem-reference.docx   Quarto reference doc
    office-templates/PPADEM-document.dotx           Word template (cover page)
    office-templates/PPADEM-document-example.docx   every style in use

Requires Quarto on PATH (or set QUARTO=/path/to/quarto) and python-docx:

    pip install python-docx
    python3 tools/build-reference-docx.py
"""

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from office_common import EU_LOGO, FONT, LOGO, OFFICE, P, ROOT, brand_theme, pandoc_default, save_reproducibly

REFERENCE = ROOT / "_extensions" / "ppadem-word" / "ppadem-reference.docx"
TEMPLATE = OFFICE / "PPADEM-document.dotx"
EXAMPLE = OFFICE / "PPADEM-document-example.docx"

# Child order of <w:pPr> required by the OOXML schema (Word rejects files
# whose elements are out of order).
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(P[name])


def set_font(style, size=None, colour=None, bold=None, italic=None):
    font = style.font
    font.name = FONT
    # Also set the East Asian / complex-script font slots so Word doesn't
    # fall back to the theme font for them.
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    if size is not None:
        font.size = Pt(size)
    if colour is not None:
        font.color.rgb = rgb(colour)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def insert_ppr(ppr, child):
    """Insert child into a <w:pPr> at its schema position, replacing any existing one."""
    name = child.tag.split("}")[1]
    old = ppr.find(qn(f"w:{name}"))
    if old is not None:
        if name != "pBdr":
            ppr.remove(old)
        else:
            # Merge border edges into the existing <w:pBdr>
            for edge in child:
                existing = old.find(edge.tag)
                if existing is not None:
                    old.remove(existing)
                old.append(edge)
            sides = ("top", "left", "bottom", "right", "between", "bar")
            for edge in sorted(old, key=lambda e: sides.index(e.tag.split("}")[1])):
                old.append(edge)
            return
    later = PPR_ORDER[PPR_ORDER.index(name) + 1:]
    for i, existing in enumerate(ppr):
        if existing.tag.split("}")[1] in later:
            ppr.insert(i, child)
            return
    ppr.append(child)


def paragraph_border(target, side: str, colour: str, size_eighths: int, space: int = 4):
    """target is a style or a paragraph."""
    ppr = target.element.get_or_add_pPr() if hasattr(target, "element") else target._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size_eighths))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), P[colour])
    borders.append(edge)
    insert_ppr(ppr, borders)


def paragraph_shading(style, colour: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), P[colour])
    insert_ppr(style.element.get_or_add_pPr(), shd)


def style_table(style):
    """Header row in soft red with a red rule; light horizontal rules."""
    el = style.element
    for old in el.findall(qn("w:tblStylePr")):
        el.remove(old)

    tblpr = el.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        el.append(tblpr)
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "insideH"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), P["ppadem-border"])
        borders.append(edge)
    later = {qn(f"w:{n}") for n in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")}
    position = next((i for i, child in enumerate(tblpr) if child.tag in later), len(tblpr))
    tblpr.insert(position, borders)

    header = OxmlElement("w:tblStylePr")
    header.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    bold = OxmlElement("w:b")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), P["ppadem-red-dark"])
    rpr.extend([bold, colour])
    tcpr = OxmlElement("w:tcPr")
    tc_borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), P["ppadem-red"])
    tc_borders.append(bottom)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), P["ppadem-red-soft"])
    tcpr.extend([tc_borders, shd])
    header.extend([rpr, tcpr])
    el.append(header)


def page_field(paragraph, code="PAGE", placeholder=None):
    """Append a field (PAGE by default) to a paragraph."""
    run = paragraph.add_run()
    parts = [("begin", None), (None, code)]
    if placeholder:
        parts.append(("separate", None))
    for kind, text in parts:
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {code} " if placeholder else code
            run._r.append(instr)
    if placeholder:
        paragraph.add_run(placeholder)
        run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    return run


# ---------------------------------------------------------------------------
# PPADEM named styles (Home → Styles in Word; custom-style="…" in Quarto)
# ---------------------------------------------------------------------------

def add_box_style(styles, name, accent, fill, text="ppadem-text", bold=False):
    """A shaded paragraph with a coloured left rule. Consecutive paragraphs in
    the same style join into one box."""
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles["Body Text"]
    style.next_paragraph_style = style
    set_font(style, size=10.5, colour=text, bold=bold)
    paragraph_shading(style, fill)
    # Same-colour top and bottom rules pad the box; Word draws them only above
    # the first and below the last paragraph of a run of identical paragraphs.
    for side in ("top", "bottom"):
        paragraph_border(style, side, fill, 4, space=6)
    paragraph_border(style, "left", accent, 24, space=8)
    fmt = style.paragraph_format
    fmt.left_indent = fmt.right_indent = Cm(0.3)
    fmt.space_before, fmt.space_after = Pt(10), Pt(10)
    # No gap between paragraphs inside the same box
    insert_ppr(style.element.get_or_add_pPr(), OxmlElement("w:contextualSpacing"))
    style.quick_style = True
    return style


def key_findings_numbering(doc, style_id: str) -> int:
    """A decimal list bound to the Key Findings style; returns its numId."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="singleLevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/>'
        f'<w:numFmt w:val="decimal"/><w:pStyle w:val="{style_id}"/><w:lvlText w:val="%1."/>'
        f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="680" w:hanging="400"/></w:pPr>'
        f'<w:rPr><w:b/><w:color w:val="{P["ppadem-teal"]}"/></w:rPr></w:lvl></w:abstractNum>'
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
    ))
    return num_id


def add_ppadem_styles(doc):
    s = doc.styles
    add_box_style(s, "PPADEM Executive Summary", "ppadem-red", "ppadem-red-soft", text="ppadem-slate")
    findings = add_box_style(s, "PPADEM Key Findings", "ppadem-teal", "ppadem-teal-soft")
    num_id = key_findings_numbering(doc, findings.style_id)
    num_pr = parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="{num_id}"/></w:numPr>')
    insert_ppr(findings.element.get_or_add_pPr(), num_pr)
    # The list indent comes from the numbering definition, not the box indent
    findings.paragraph_format.left_indent = None
    add_box_style(s, "PPADEM Note", "ppadem-teal", "ppadem-teal-soft")
    add_box_style(s, "PPADEM Warning", "ppadem-gold", "ppadem-gold-soft")
    add_box_style(s, "PPADEM Important", "ppadem-red", "ppadem-red-soft")

    quote = s.add_style("PPADEM Pull Quote", WD_STYLE_TYPE.PARAGRAPH)
    quote.base_style = s["Body Text"]
    set_font(quote, size=14, colour="ppadem-slate", italic=True)
    paragraph_border(quote, "left", "ppadem-red", 32, space=12)
    quote.paragraph_format.left_indent = Cm(0.6)
    quote.paragraph_format.space_before = Pt(12)
    quote.paragraph_format.space_after = Pt(2)
    quote.paragraph_format.keep_with_next = True
    quote.quick_style = True

    attribution = s.add_style("PPADEM Quote Attribution", WD_STYLE_TYPE.PARAGRAPH)
    attribution.base_style = s["Body Text"]
    set_font(attribution, size=9.5, colour="ppadem-gray", bold=True)
    attribution.paragraph_format.left_indent = Cm(0.6)
    attribution.paragraph_format.space_after = Pt(12)
    attribution.next_paragraph_style = s["Body Text"]
    quote.next_paragraph_style = attribution

    label = s.add_style("PPADEM Label", WD_STYLE_TYPE.PARAGRAPH)
    label.base_style = s["Body Text"]
    set_font(label, size=8.5, colour="ppadem-red", bold=True)
    label.font.all_caps = True
    label.paragraph_format.space_before = Pt(10)
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True
    label.quick_style = True


# ---------------------------------------------------------------------------
# Styles and page setup shared by all three files
# ---------------------------------------------------------------------------

def build_styles(doc):
    s = doc.styles

    # Body text
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        set_font(s[name], size=10.5, colour="ppadem-text")
    s["Body Text"].paragraph_format.space_after = Pt(6)
    s["Body Text"].paragraph_format.line_spacing = 1.15

    # Title block
    set_font(s["Title"], size=26, colour="ppadem-red-dark", bold=True)
    s["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s["Title"].paragraph_format.space_after = Pt(4)
    set_font(s["Subtitle"], size=14, colour="ppadem-teal", bold=False, italic=False)
    s["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name in ("Author", "Date"):
        set_font(s[name], size=10, colour="ppadem-gray")
        s[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(s["Abstract Title"], size=9, colour="ppadem-red", bold=True)
    set_font(s["Abstract"], size=10.5, colour="ppadem-slate")
    paragraph_shading(s["Abstract"], "ppadem-red-soft")
    paragraph_border(s["Abstract"], "left", "ppadem-red", 24, space=8)

    # Headings: H1 gets the PPADEM underline
    set_font(s["Heading 1"], size=16, colour="ppadem-red-dark", bold=True)
    s["Heading 1"].paragraph_format.space_before = Pt(18)
    s["Heading 1"].paragraph_format.space_after = Pt(8)
    paragraph_border(s["Heading 1"], "bottom", "ppadem-red-border", 8)
    set_font(s["Heading 2"], size=13, colour="ppadem-slate", bold=True)
    s["Heading 2"].paragraph_format.space_before = Pt(14)
    for level in range(3, 7):
        set_font(s[f"Heading {level}"], size=11, colour="ppadem-slate", bold=True, italic=False)
    for level in range(1, 7):
        set_font(s[f"Heading {level} Char"], colour="ppadem-red-dark" if level == 1 else "ppadem-slate")
    set_font(s["TOC Heading"], size=16, colour="ppadem-red-dark", bold=True)

    # Pressing Enter after a heading or the title block continues in body text
    for name in [f"Heading {level}" for level in range(1, 10)] + ["Title", "Subtitle", "Author", "Date"]:
        s[name].next_paragraph_style = s["Body Text"]

    # Quotes, captions, links
    set_font(s["Block Text"], size=11, colour="ppadem-slate", italic=False)
    paragraph_border(s["Block Text"], "left", "ppadem-teal", 24, space=10)
    paragraph_shading(s["Block Text"], "ppadem-teal-soft")
    for name in ("Caption", "Table Caption", "Image Caption"):
        set_font(s[name], size=9, colour="ppadem-gray", italic=False)
    s["Hyperlink"].font.color.rgb = rgb("ppadem-red")
    style_table(s["Table"])

    add_ppadem_styles(doc)

    # Theme colours and fonts: the colour picker shows the PPADEM palette
    theme = doc.part.part_related_by(RT.THEME)
    theme._blob = brand_theme(theme.blob)


def page_setup(doc, cover=False):
    """A4, logo in the header, "PPADEM Project · Page n" in the footer.

    With cover=True the first page has no header or footer (the cover page
    carries its own large logo)."""
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.2)
        # Word requires these on <w:pgMar>; Pandoc's default omits them
        section.header_distance = section.footer_distance = Cm(1.25)
        section.gutter = Cm(0)
        section.different_first_page_header_footer = cover

        head = section.header.paragraphs[0]
        head.text = ""
        head.add_run().add_picture(str(LOGO), height=Cm(1.0))

        foot = section.footer.paragraphs[0]
        foot.text = ""
        run = foot.add_run("PPADEM Project  ·  Page ")
        run.font.size, run.font.color.rgb = Pt(8), rgb("ppadem-gray")
        page = page_field(foot)
        page.font.size, page.font.color.rgb = Pt(8), rgb("ppadem-gray")


def build_base():
    with tempfile.TemporaryDirectory() as tmp:
        doc = Document(str(pandoc_default("reference.docx", Path(tmp) / "reference.docx")))
    build_styles(doc)
    return doc


# ---------------------------------------------------------------------------
# Template and example content
# ---------------------------------------------------------------------------

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def brand_rule(doc):
    """Short red / gold / teal rule, as under the slide and report titles."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for column in table.columns:
        column.width = Cm(1.3)
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    width.set(qn("w:w"), str(int(Cm(3.9).twips)))
    width.set(qn("w:type"), "dxa")
    tbl_pr.find(qn("w:jc")).addnext(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>'))
    for cell, fill in zip(table.rows[0].cells, ("ppadem-red", "ppadem-gold", "ppadem-teal")):
        cell.width = Cm(1.3)
        cell._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{P[fill]}"/>')
        )
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = Pt(1)
        para.add_run().font.size = Pt(1)
    row = table.rows[0]._tr.get_or_add_trPr()
    row.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="80" w:hRule="exact"/>'))


def cover(doc, label, title, subtitle, author, date):
    doc.add_paragraph().add_run().add_picture(str(LOGO), height=Cm(2.6))
    lead = doc.add_paragraph(label, style="PPADEM Label")
    lead.paragraph_format.space_before = Pt(150)
    block = [
        doc.add_paragraph(title, style="Title"),
        doc.add_paragraph(subtitle, style="Subtitle"),
        doc.add_paragraph(author, style="Author"),
        doc.add_paragraph(date, style="Date"),
    ]
    # A red bar down the left of the title block, as on the title slide
    for para in block:
        paragraph_border(para, "left", "ppadem-red", 48, space=12)
    block[-1].paragraph_format.space_after = Pt(18)
    brand_rule(doc)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Contents", style="TOC Heading")
    toc = doc.add_paragraph(style="Body Text")
    page_field(toc, 'TOC \\o "1-3" \\h \\z \\u',
               "Right-click here and choose Update Field to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def funding(doc):
    """The EU funding logo, centred at the end of the document."""
    para = doc.add_paragraph(style="Body Text")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(24)
    para.add_run().add_picture(str(EU_LOGO), width=Cm(7.0))


def build_template(doc):
    clear_body(doc)
    cover(doc, "[Document type]", "[Document title]", "[Subtitle]", "[Authors]", "[Date]")
    doc.add_paragraph("[Introduction]", style="Heading 1")
    doc.add_paragraph(
        "[Start writing here. Apply PPADEM styles from Home → Styles: headings, "
        "PPADEM Note, PPADEM Key Findings and more.]",
        style="Body Text",
    )
    funding(doc)
    page_setup(doc, cover=True)


def say(doc, text, style="Body Text"):
    return doc.add_paragraph(text, style=style)


def build_example(doc):
    clear_body(doc)
    cover(doc, "Guide", "Using the PPADEM Word template",
          "How to write branded documents with Word styles", "PPADEM Project Team", "October 2026")

    say(doc, "Getting started", "Heading 1")
    say(doc, "Everything in this document is formatted with styles, so it stays on brand "
             "without any manual formatting. Click in a paragraph and choose a style from "
             "Home → Styles. To see every style, open the Styles pane (the small arrow at the "
             "bottom-right of the Styles group on Windows, or Format → Style on a Mac).")
    say(doc, "To start your own document, double-click PPADEM-document.dotx. Word opens a new "
             "untitled copy, so the template itself is never changed.")
    say(doc, "PPADEM Executive Summary", "PPADEM Label")
    say(doc, "This box uses the PPADEM Executive Summary style. Use it for a short summary at "
             "the start of a report.", "PPADEM Executive Summary")
    say(doc, "Consecutive paragraphs in the same box style join into one box.",
        "PPADEM Executive Summary")

    say(doc, "Headings", "Heading 1")
    say(doc, "Use Heading 1 for main sections. It gets the red rule underneath and appears in "
             "the table of contents.")
    say(doc, "Heading 2 for subsections", "Heading 2")
    say(doc, "Heading 2 and Heading 3 are slate grey. Press Enter after a heading and Word "
             "switches back to Body Text automatically.")
    say(doc, "Heading 3 for smaller divisions", "Heading 3")
    say(doc, "After adding or renaming headings, right-click the table of contents and choose "
             "Update Field.")

    say(doc, "Boxes and callouts", "Heading 1")
    say(doc, "Key findings", "PPADEM Label")
    for finding in (
        "Use the PPADEM Key Findings style for a numbered list of findings.",
        "Each paragraph is numbered automatically.",
        "To restart numbering in a new box, right-click the number and choose Restart at 1.",
    ):
        say(doc, finding, "PPADEM Key Findings")
    say(doc, "PPADEM Note: teal, for background information or tips.", "PPADEM Note")
    say(doc, "PPADEM Warning: gold, for caveats and things to watch out for.", "PPADEM Warning")
    say(doc, "PPADEM Important: red, for the one thing a reader must not miss.", "PPADEM Important")
    say(doc, "PPADEM Pull Quote: “Parties are where democracy is organised, or where it fails.”",
        "PPADEM Pull Quote")
    say(doc, "— PPADEM Quote Attribution", "PPADEM Quote Attribution")
    say(doc, "Block Text is a quieter teal quote style for longer quotations from interviews or "
             "documents.", "Block Text")

    say(doc, "Tables and figures", "Heading 1")
    say(doc, "Insert a table, then choose the PPADEM “Table” style from Table Design → Table Styles "
             "(it is under Custom at the top). Put a Table Caption paragraph above it.")
    say(doc, "Table 1: Coverage by region", "Table Caption")
    rows = [("Region", "Countries", "Parties"), ("West Africa", "9", "30"),
            ("East Africa", "7", "22"), ("Southern Africa", "9", "28")]
    table = doc.add_table(rows=len(rows), cols=3, style="Table")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
            if c:
                table.cell(r, c).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    say(doc, "Figure captions go below the figure in the Image Caption style.", "Image Caption")

    say(doc, "Colours", "Heading 1")
    say(doc, "When you colour text, shapes or charts, pick from the Theme Colours row at the top "
             "of the colour menu: it holds the PPADEM red, teal, gold and grey. Charts inserted in "
             "Word use these colours automatically.")
    say(doc, "Funding logo", "Heading 1")
    say(doc, "End every document with the EU funding logo, as below. Keep it on the last page.")
    funding(doc)
    page_setup(doc, cover=True)


def main():
    reference = build_base()
    page_setup(reference)
    save_reproducibly(reference, REFERENCE)

    template = build_base()
    build_template(template)
    save_reproducibly(template, TEMPLATE, template=True)

    example = build_base()
    build_example(example)
    save_reproducibly(example, EXAMPLE)


if __name__ == "__main__":
    main()
